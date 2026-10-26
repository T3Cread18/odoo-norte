import base64
import csv
import datetime
import io
import json
import logging
import re
import unicodedata
from datetime import timedelta

from markupsafe import Markup

from odoo import api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Campos a los que se puede mapear una columna del archivo, en el orden en
# que se muestran en la vista previa.
CAMPOS = [
    ("cedula", "Cédula / RIF"),
    ("nombre", "Nombre / razón social"),
    ("contratista", "Contratista"),
    ("vendedor", "Vendedor"),
    ("tipo_servicio", "Tipo de servicio"),
    ("telefono_movil", "Teléfono móvil"),
    ("telefono_fijo", "Teléfono fijo"),
    ("direccion", "Dirección"),
    ("franquicia", "Franquicia / sede"),
    ("zona", "Zona"),
    ("antiguedad", "Antigüedad (días)"),
    ("fecha_asignacion", "Fecha de asignación"),
    ("estatus", "Estatus de la orden (sistema)"),
    ("fecha_finalizacion", "Fecha de finalización"),
]
ETIQUETAS = dict(CAMPOS)
OBLIGATORIOS = ("cedula", "nombre")

# Encabezados que se reconocen solos (normalizados: mayúsculas, sin acentos
# ni puntos). El mapeo que el usuario elige se recuerda en PARAM_MAPEO y
# tiene prioridad sobre estos.
ALIAS = {
    # El primer alias gana si el archivo trae varios (ver _sugerencia)
    "tipo_servicio": ["TIPO SERVICIO", "TIPO DE SERVICIO", "SERVICIO"],
    "cedula": ["RIF", "CEDULA", "CI", "RIF CEDULA", "CEDULA RIF"],
    "nombre": ["RAZON SOCIAL", "NOMBRE", "CLIENTE", "NOMBRE CLIENTE"],
    "telefono_movil": ["TLF MOVIL", "TELEFONO MOVIL", "MOVIL", "CELULAR"],
    "telefono_fijo": ["TLF FIJO", "TELEFONO FIJO", "FIJO"],
    "antiguedad": ["ANTIGUEDAD", "DIAS"],
    "contratista": ["TECNICO", "CONTRATISTA"],
    "vendedor": ["VENDEDOR", "VENDEDORA", "ASESOR", "ASESOR DE VENTAS"],
    "direccion": ["REFERENCIA", "DIRECCION"],
    "franquicia": ["FRANQUICIA", "SEDE"],
    # En el export del sistema ZONA es la región (siempre ACARIGUA): se
    # prefiere SECTOR; si no coincide con una zona, se busca en la dirección
    "zona": ["SECTOR", "ZONA"],
    "fecha_asignacion": ["FECHA ASIGNACION", "FECHA DE ASIGNACION"],
    "estatus": ["ESTATUS ORDEN SERVICIO", "ESTATUS ORDEN", "ESTATUS"],
    "fecha_finalizacion": ["FECHA FINALIZACION", "FECHA DE FINALIZACION"],
}
PARAM_MAPEO = "atc_visitas_tecnicas.import_instalaciones_mapeo"
# Campo de la orden que actualiza cada dato del archivo cuando la orden ya
# existe (abierta, misma cédula).
ACTUALIZABLES = {
    "tipo_servicio": "tipo_servicio_id",
    "contratista": "contratista_id",
    "vendedor": "vendedor_id",
    "telefono_movil": "telefono_movil",
    "telefono_fijo": "telefono_fijo",
    "direccion": "direccion",
    "franquicia": "franquicia_id",
}

ESTADOS_ABIERTOS = ("pendiente", "asignada")
FILAS_PREVIA = 5

# Palabras que delatan una dirección: si aparecen en la columna del
# contratista, lo más probable es que las columnas estén cruzadas.
PALABRAS_DIRECCION = {
    "CALLE", "CALLEJON", "AV", "AVENIDA", "CARRERA", "CRA", "TRANSVERSAL",
    "DIAGONAL", "BARRIO", "URB", "URBANIZACION", "SECTOR", "VEREDA", "CASA",
    "QTA", "QUINTA", "ESQUINA", "EDIF", "EDIFICIO", "RESIDENCIAS", "CONJUNTO",
    "MANZANA", "PARCELA", "KM", "CARRETERA", "CASERIO", "FRENTE", "DETRAS",
}
# Celdas que el sistema de origen llena cuando no hay dato (p. ej. TECNICO
# "N/A" = aún sin contratista): se tratan como vacías.
SIN_DATO = {"N/A", "NA", "N/D", "-", "--"}
# Cédula ya normalizada: solo números, o RIF J/G
CEDULA_VALIDA = re.compile(r"^(\d{5,9}|[JG]-\d{8}(-\d)?)$")


def _sin_acentos(txt):
    return "".join(c for c in unicodedata.normalize("NFD", txt)
                   if unicodedata.category(c) != "Mn")


def _norm(val):
    if val is None:
        return ""
    txt = _sin_acentos(str(val)).upper().replace(".", " ")
    return re.sub(r"\s+", " ", txt).strip()


def _limpiar(val):
    """Texto de una celda sin espacios repetidos."""
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    return re.sub(r"\s+", " ", str(val)).strip()


def normalizar_cedula(rif):
    """'V-16292774-0' / 'V-09250618' -> '16292774' / '9250618' (formato de
    res.partner.cedula). Los RIF jurídicos (J/G) se dejan como 'J-40880242-5'."""
    txt = _limpiar(rif).upper().replace(" ", "")
    m = re.match(r"^([VEJGP])?-?0*(\d+)(?:-(\d))?$", txt)
    if not m:
        return txt
    letra, numero, digito = m.groups()
    if letra in ("J", "G"):
        return "%s-%s%s" % (letra, numero.zfill(8),
                            "-%s" % digito if digito else "")
    return numero


def normalizar_telefono(val):
    """'4121880078' -> '04121880078'."""
    if val is None or val == "":
        return False
    if isinstance(val, float):
        val = int(val)
    digitos = re.sub(r"\D", "", str(val))
    if not digitos:
        return False
    if len(digitos) == 10 and not digitos.startswith("0"):
        digitos = "0" + digitos
    return digitos


def estatus_origen(val):
    """Estatus de la orden en el sistema de origen: 'finalizada',
    'cancelada' o 'abierta' (POR ACTIVAR, PENDIENTE, ASIGNADA, vacío...)."""
    txt = _norm(val)
    if txt.startswith("FINALIZ"):
        return "finalizada"
    if txt.startswith(("CANCEL", "ANUL")):
        return "cancelada"
    return "abierta"


def leer_fecha(val):
    """Fecha de una celda: ISO (así se guardan las fechas de Excel), o
    dd/mm/aaaa, dd-mm-aaaa, dd/mm/aa. None si no hay o no se entiende."""
    txt = _limpiar(val)
    if not txt:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.datetime.strptime(txt[:10], fmt).date()
        except ValueError:
            continue
    return None


def _celda_json(val):
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.isoformat()[:10]
    if isinstance(val, (int, float, str)) or val is None:
        return val
    return str(val)


class AtcInstalacionImport(models.TransientModel):
    _name = "atc.instalacion.import"
    _description = "Importar instalaciones pendientes desde Excel"

    archivo = fields.Binary(string="Archivo (.xlsx, .xls o .csv)", required=True)
    nombre_archivo = fields.Char(string="Nombre del archivo")
    fecha_reporte = fields.Date(
        string="Fecha del reporte", required=True,
        default=fields.Date.context_today,
        help="Día en que se generó el reporte. Si no hay columna de fecha de "
             "asignación, se calcula como esta fecha menos la ANTIGÜEDAD.")
    state = fields.Selection(
        [("input", "Archivo"), ("mapeo", "Columnas"), ("done", "Resultado")],
        default="input")
    hoja = fields.Char(string="Hoja", readonly=True)
    fila_encabezado = fields.Integer(
        string="Fila de encabezados",
        help="Número de fila (como en Excel) donde están los títulos de las "
             "columnas. Se detecta solo; cámbiala y pulsa 'Volver a leer' si "
             "no es la correcta.")
    datos = fields.Text(help="Filas del archivo bajo el encabezado, en JSON.")
    total_filas = fields.Integer(string="Filas con datos", readonly=True)
    linea_ids = fields.One2many(
        "atc.instalacion.import.columna", "wizard_id", string="Columnas")
    vista_previa = fields.Html(
        string="Vista previa", compute="_compute_vista_previa",
        sanitize=False)
    revision = fields.Html(
        string="Revisión del archivo", compute="_compute_revision",
        sanitize=False)
    hay_errores = fields.Boolean(compute="_compute_revision")
    omitir_errores = fields.Boolean(
        string="Importar sin las filas con error",
        help="Importa el resto del archivo y deja fuera las filas con error, "
             "que se listan en el resultado para cargarlas después.")
    resultado = fields.Text(string="Resultado", readonly=True)

    # ------------------------------------------------------------------
    # Lectura del archivo
    # ------------------------------------------------------------------
    def _leer_hojas(self):
        """[(nombre_hoja, filas)] con los valores crudos de cada celda."""
        data = base64.b64decode(self.archivo)
        nombre = (self.nombre_archivo or "").lower()
        try:
            if nombre.endswith(".csv"):
                return [("CSV", self._leer_csv(data))]
            if nombre.endswith(".xls"):
                return self._leer_xls(data)
            return self._leer_xlsx(data)
        except UserError:
            raise
        except Exception as exc:
            raise UserError(self.env._(
                "No se pudo leer el archivo: %s") % exc)

    def _leer_xlsx(self, data):
        import openpyxl
        wb = openpyxl.load_workbook(
            io.BytesIO(data), data_only=True, read_only=True)
        hojas = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                 for ws in wb.worksheets]
        wb.close()
        return hojas

    def _leer_xls(self, data):
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        hojas = []
        for sh in wb.sheets():
            filas = []
            for r in range(sh.nrows):
                fila = []
                for c in range(sh.ncols):
                    celda = sh.cell(r, c)
                    if celda.ctype == xlrd.XL_CELL_DATE:
                        fila.append(xlrd.xldate_as_datetime(
                            celda.value, wb.datemode))
                    else:
                        fila.append(celda.value if celda.value != "" else None)
                filas.append(fila)
            hojas.append((sh.name, filas))
        return hojas

    def _leer_csv(self, data):
        try:
            texto = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = data.decode("latin-1")
        try:
            delimitador = csv.Sniffer().sniff(
                texto[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            # El Sniffer falla con filas largas o comas dentro de los campos
            # (export del sistema, separado por ';'): el más frecuente en
            # la primera línea
            primera = texto.split("\n", 1)[0]
            delimitador = max(",;\t", key=primera.count)
        return [[c if c != "" else None for c in fila]
                for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]

    # ------------------------------------------------------------------
    # Mapeo de columnas
    # ------------------------------------------------------------------
    def _mapeo_recordado(self):
        try:
            return json.loads(self.env["ir.config_parameter"].sudo().get_param(
                PARAM_MAPEO) or "{}")
        except ValueError:
            return {}

    def _sugerir_campo(self, encabezado, recordado):
        return self._sugerencia(encabezado, recordado)[0]

    def _sugerencia(self, encabezado, recordado):
        """(campo, prioridad); menor prioridad = mejor. Gana el mapeo
        recordado y, entre alias, el que está antes en la lista: si el
        archivo trae SERVICIO y TIPO SERVICIO, se usa TIPO SERVICIO."""
        clave = _norm(encabezado)
        if not clave:
            return False, 0
        if clave in recordado:
            return recordado[clave] or False, -1
        for campo, alias in ALIAS.items():
            if clave in alias:
                return campo, alias.index(clave)
        return False, 0

    def _detectar_encabezado(self, hojas, recordado):
        """(hoja, índice de fila) del encabezado: la primera fila, en las 20
        primeras de cada hoja, con más títulos reconocidos (mínimo 2). Si no
        hay, la primera fila no vacía de la primera hoja."""
        mejor = None
        for nombre, filas in hojas:
            for i, fila in enumerate(filas[:20]):
                puntos = sum(1 for c in fila
                             if self._sugerir_campo(c, recordado))
                if puntos >= 2 and (mejor is None or puntos > mejor[0]):
                    mejor = (puntos, nombre, i)
        if mejor:
            return mejor[1], mejor[2]
        nombre, filas = hojas[0]
        for i, fila in enumerate(filas):
            if any(_limpiar(c) for c in fila):
                return nombre, i
        raise UserError(self.env._("El archivo está vacío."))

    def _cargar(self, hoja=None, indice=None):
        """Lee el archivo y prepara las líneas de mapeo. Sin hoja/índice,
        detecta el encabezado."""
        self.ensure_one()
        if not self.archivo:
            raise UserError(self.env._("Debes adjuntar un archivo."))
        hojas = self._leer_hojas()
        recordado = self._mapeo_recordado()
        if hoja is None:
            hoja, indice = self._detectar_encabezado(hojas, recordado)
        filas = dict(hojas).get(hoja) or hojas[0][1]
        if indice >= len(filas):
            raise UserError(self.env._(
                "La hoja solo tiene %s filas.") % len(filas))
        encabezado = filas[indice]
        cuerpo = [f for f in filas[indice + 1:] if any(_limpiar(c) for c in f)]
        ancho = max([len(encabezado)] + [len(f) for f in cuerpo])

        columnas, mejor = [], {}   # mejor: {campo: (prioridad, col)}
        for col in range(ancho):
            titulo = _limpiar(encabezado[col]) if col < len(encabezado) else ""
            valores = [_limpiar(f[col]) for f in cuerpo if col < len(f)]
            valores = [v for v in valores if v]
            if not titulo and not valores:
                continue
            campo, prioridad = self._sugerencia(titulo, recordado)
            if campo and (campo not in mejor or prioridad < mejor[campo][0]):
                mejor[campo] = (prioridad, col)
            columnas.append((col, titulo, valores, campo))
        # Cada campo, a una sola columna: la de mejor prioridad
        lineas = [(0, 0, {
            "indice": col,
            "columna": titulo or self.env._("(columna %s sin título)") % (col + 1),
            "ejemplo": " | ".join(valores[:3])[:120],
            "campo": campo if campo and mejor[campo][1] == col else False,
        }) for col, titulo, valores, campo in columnas]
        self.linea_ids.unlink()
        self.write({
            "state": "mapeo",
            "hoja": hoja,
            "fila_encabezado": indice + 1,
            "datos": json.dumps([[_celda_json(c) for c in f] for f in cuerpo]),
            "total_filas": len(cuerpo),
            "linea_ids": lineas,
            "resultado": False,
        })

    def _mapeo(self):
        """{campo: índice de columna}; valida obligatorios y duplicados."""
        mapeo = {}
        for linea in self.linea_ids.filtered("campo"):
            if linea.campo in mapeo:
                otra = self.linea_ids.filtered(
                    lambda l: l.indice == mapeo[linea.campo])
                raise UserError(self.env._(
                    "Las columnas «%(a)s» y «%(b)s» están mapeadas al mismo "
                    "campo «%(campo)s».", a=otra.columna, b=linea.columna,
                    campo=ETIQUETAS[linea.campo]))
            mapeo[linea.campo] = linea.indice
        faltan = [ETIQUETAS[c] for c in OBLIGATORIOS if c not in mapeo]
        if faltan:
            raise UserError(self.env._(
                "Falta mapear: %s.") % ", ".join(faltan))
        return mapeo

    def _filas(self, mapeo, limite=None):
        """[(nº de fila en el archivo, dict normalizado)]."""
        cuerpo = json.loads(self.datos or "[]")
        primera = self.fila_encabezado + 1
        res = []
        for n, f in enumerate(cuerpo[:limite] if limite else cuerpo):
            def get(campo):
                i = mapeo.get(campo)
                val = f[i] if i is not None and i < len(f) else None
                return None if _norm(val) in SIN_DATO else val

            cedula = normalizar_cedula(get("cedula")) if get("cedula") else ""
            nombre = _limpiar(get("nombre"))
            if not cedula and not nombre:
                continue
            antig = re.search(r"\d+", _limpiar(get("antiguedad")))
            res.append((primera + n, {
                "tipo_servicio": _limpiar(get("tipo_servicio")),
                "cedula": cedula,
                "nombre": nombre,
                "telefono_movil": normalizar_telefono(get("telefono_movil")),
                "telefono_fijo": normalizar_telefono(get("telefono_fijo")),
                "antiguedad": int(antig.group()) if antig else None,
                "fecha_asignacion": leer_fecha(get("fecha_asignacion")),
                "contratista": _limpiar(get("contratista")),
                "vendedor": _limpiar(get("vendedor")),
                "direccion": _limpiar(get("direccion")),
                "franquicia": _limpiar(get("franquicia")),
                "zona": _limpiar(get("zona")),
                "estatus": estatus_origen(get("estatus")),
                "fecha_finalizacion": leer_fecha(get("fecha_finalizacion")),
            }))
        return res

    def _fecha_asignacion(self, fila):
        if fila["fecha_asignacion"]:
            return fila["fecha_asignacion"]
        return self.fecha_reporte - timedelta(days=fila["antiguedad"] or 0)

    def _revisar(self, mapeo, filas):
        """Problemas del archivo antes de importar.

        Devuelve ({nº fila: [errores]}, [avisos]). Una fila con error no se
        debe importar tal cual (datos en la columna equivocada, cédula que no
        se entiende); los avisos solo piden revisar algo."""
        _ = self.env._
        Contratista = self.env["atc.contratista"].with_context(active_test=False)
        Servicio = self.env["atc.tipo.servicio"].with_context(active_test=False)
        contratistas = Contratista.search([]).mapped("name")
        # Igual que _buscar_o_crear (=ilike): sin distinguir mayúsculas
        existentes_c = {n.lower() for n in contratistas}
        contratistas_norm = {_norm(n) for n in contratistas}
        existentes_s = {n.lower() for n in Servicio.search([]).mapped("name")}
        existentes_f = {n.lower() for n in self.env["atc.franquicia"].with_context(
            active_test=False).search([]).mapped("name")}
        existentes_v = {n.lower() for n in self.env["atc.vendedor"].with_context(
            active_test=False).search([]).mapped("name")}

        errores = {}
        nuevos_c, nuevos_s, nuevas_f, nuevos_v = {}, {}, {}, {}
        sin_cedula, sin_contratista, telefonos, sin_antiguedad = [], [], [], []
        for n, f in filas:
            problemas = []
            contratista, direccion = f["contratista"], f["direccion"]
            nuevo = contratista and contratista.lower() not in existentes_c
            palabras = set(re.sub(r"[^A-Z0-9]+", " ", _norm(contratista)).split())
            if (nuevo or not contratista) and direccion and (
                    _norm(direccion) in contratistas_norm
                    or (nuevo and palabras & PALABRAS_DIRECCION)):
                problemas.append(_(
                    "las columnas de contratista y dirección parecen "
                    "intercambiadas (contratista: «%(c)s», dirección: «%(d)s»)",
                    c=contratista or _("vacío"), d=direccion))
            elif nuevo and palabras & PALABRAS_DIRECCION:
                problemas.append(_(
                    "el contratista «%s» parece una dirección", contratista))
            if f["cedula"] and not CEDULA_VALIDA.match(f["cedula"]):
                problemas.append(_("cédula/RIF no válido «%s»", f["cedula"]))
            if not f["nombre"]:
                problemas.append(_("falta el nombre del cliente"))
            if problemas:
                errores[n] = problemas
                continue
            if nuevo:
                nuevos_c.setdefault(contratista, []).append(n)
            elif not contratista:
                sin_contratista.append(n)
            servicio = f["tipo_servicio"]
            if servicio and servicio.lower() not in existentes_s:
                nuevos_s.setdefault(servicio, []).append(n)
            franquicia = f["franquicia"]
            if franquicia and franquicia.lower() not in existentes_f:
                nuevas_f.setdefault(franquicia, []).append(n)
            vendedor = f["vendedor"]
            if vendedor and vendedor.lower() not in existentes_v:
                nuevos_v.setdefault(vendedor, []).append(n)
            if not f["cedula"]:
                sin_cedula.append(n)
            for tel in (f["telefono_movil"], f["telefono_fijo"]):
                if tel and len(tel) != 11:
                    telefonos.append("%s (%s)" % (n, tel))
            if ("antiguedad" in mapeo and "fecha_asignacion" not in mapeo
                    and f["antiguedad"] is None):
                sin_antiguedad.append(n)

        def filas_txt(nums):
            return ", ".join(map(str, nums))

        avisos = []
        for nombre, nums in nuevos_c.items():
            avisos.append(_(
                "Contratista nuevo «%(c)s» (filas %(f)s): no existe en Odoo; se "
                "creará y no verá sus órdenes en el portal hasta que se le cree "
                "un usuario. Si es un contratista que ya existe, escríbelo igual "
                "que en Odoo.", c=nombre, f=filas_txt(nums)))
        for nombre, nums in nuevos_s.items():
            avisos.append(_("Tipo de servicio nuevo «%(s)s» (filas %(f)s): se "
                            "creará.", s=nombre, f=filas_txt(nums)))
        for nombre, nums in nuevas_f.items():
            avisos.append(_("Franquicia nueva «%(s)s» (filas %(f)s): se creará; "
                            "revisa en Configuración → Franquicias qué zonas "
                            "atiende.", s=nombre, f=filas_txt(nums)))
        if nuevos_v:
            avisos.append(_(
                "Vendedores nuevos (se crearán; no verán sus ventas en el "
                "portal hasta que se les cree un usuario): %s.",
                "; ".join("%s (%s)" % (v, len(nums))
                          for v, nums in nuevos_v.items())))
        if sin_contratista:
            avisos.append(_("Sin contratista, quedarán pendientes de asignar: "
                            "filas %s.", filas_txt(sin_contratista)))
        if sin_cedula:
            avisos.append(_("Sin cédula/RIF (no se puede saber si ya existen, se "
                            "crean siempre): filas %s.", filas_txt(sin_cedula)))
        if telefonos:
            avisos.append(_("Teléfonos que no tienen 11 dígitos: filas %s.",
                            ", ".join(telefonos)))
        if sin_antiguedad:
            avisos.append(_("Antigüedad vacía o sin número, se toma la fecha del "
                            "reporte como asignación: filas %s.",
                            filas_txt(sin_antiguedad)))
        return errores, avisos

    @staticmethod
    def _errores_txt(errores):
        return ["Fila %s: %s" % (n, "; ".join(p))
                for n, p in sorted(errores.items())]

    @api.depends("linea_ids.campo", "datos", "fila_encabezado")
    def _compute_revision(self):
        for w in self:
            w.revision = False
            w.hay_errores = False
            if w.state != "mapeo" or not w.datos:
                continue
            mapeo = {l.campo: l.indice for l in w.linea_ids if l.campo}
            if not all(c in mapeo for c in OBLIGATORIOS):
                continue
            errores, avisos = w._revisar(mapeo, w._filas(mapeo))
            w.hay_errores = bool(errores)
            html = Markup("")
            if errores:
                html += Markup(
                    '<div class="alert alert-danger" role="alert"><b>%s</b>'
                    "<ul>%s</ul>%s</div>") % (
                    w.env._("%s fila(s) con error: corrige el archivo y vuelve "
                            "a cargarlo (botón «Cambiar archivo»).", len(errores)),
                    Markup("").join(Markup("<li>%s</li>") % e
                                    for e in w._errores_txt(errores)),
                    w.env._("Si no puedes corregirlo ahora, marca «Importar sin "
                            "las filas con error» para cargar el resto."))
            if avisos:
                html += Markup(
                    '<div class="alert alert-warning" role="alert"><b>%s</b>'
                    "<ul>%s</ul></div>") % (
                    w.env._("Revisa antes de importar:"),
                    Markup("").join(Markup("<li>%s</li>") % a for a in avisos))
            if not errores and not avisos:
                html = Markup('<div class="alert alert-success" role="alert">%s</div>') \
                    % w.env._("Revisión del archivo: sin problemas.")
            w.revision = html

    @api.depends("linea_ids.campo", "datos", "fecha_reporte",
                 "fila_encabezado")
    def _compute_vista_previa(self):
        for w in self:
            w.vista_previa = False
            if w.state != "mapeo" or not w.datos:
                continue
            mapeo = {l.campo: l.indice for l in w.linea_ids if l.campo}
            if not mapeo:
                continue
            # Antigüedad y fecha se muestran juntas como la fecha de
            # asignación resultante.
            campos = [c for c, _ in CAMPOS if c in mapeo
                      and c not in ("antiguedad", "fecha_asignacion")]
            cab = Markup("").join(
                Markup("<th>%s</th>") % ETIQUETAS[c] for c in campos)
            cab += Markup("<th>%s</th>") % w.env._("Asignada el")
            cuerpo = Markup("")
            for _n, fila in w._filas(mapeo, limite=FILAS_PREVIA):
                celdas = Markup("").join(
                    Markup("<td>%s</td>") % (fila[c] or "") for c in campos)
                fecha = w._fecha_asignacion(fila) if w.fecha_reporte else ""
                celdas += Markup("<td>%s</td>") % (fecha or "")
                cuerpo += Markup("<tr>%s</tr>") % celdas
            w.vista_previa = Markup(
                '<table class="table table-sm table-striped o_list_table">'
                "<thead><tr>%s</tr></thead><tbody>%s</tbody></table>"
            ) % (cab, cuerpo)

    # ------------------------------------------------------------------
    # Catálogos
    # ------------------------------------------------------------------
    def _buscar_o_crear(self, model, name, cache):
        if not name:
            return model.browse()
        clave = _norm(name)
        if clave not in cache:
            rec = model.with_context(active_test=False).search(
                [("name", "=ilike", name)], limit=1)
            if not rec:
                rec = model.create({"name": name})
            elif not rec.active:
                rec.active = True
            cache[clave] = rec
        return cache[clave]

    def _buscar_franquicia(self, name, cache):
        # Siempre una franquicia hija (el padre ACARIGUA se llama igual que una hija)
        if not name:
            return self.env["atc.franquicia"]
        clave = _norm(name)
        if clave not in cache:
            cache[clave] = self.env["atc.franquicia"]._buscar_hija(name)
        return cache[clave]

    def _zona_por_direccion(self, direccion, zonas):
        """Primera zona cuyo nombre aparece como palabra completa en la
        dirección; se prueban primero los nombres más largos para que
        'VILLA ARAURE' gane sobre 'ARAURE'."""
        texto = " %s " % re.sub(r"[^A-Z0-9]+", " ", _norm(direccion))
        for clave, zona in zonas:
            if " %s " % clave in texto:
                return zona
        return self.env["atc.zona"]

    # ------------------------------------------------------------------
    # Importación
    # ------------------------------------------------------------------
    def _importar(self, omitir_errores=False):
        """Crea/actualiza las instalaciones y devuelve el resumen. Las filas
        con error (ver _revisar) bloquean la importación salvo que se pida
        omitirlas."""
        mapeo = self._mapeo()
        filas = self._filas(mapeo)
        if not filas:
            raise UserError(self.env._("No hay filas con cédula o nombre."))
        con_error, avisos = self._revisar(mapeo, filas)
        if con_error and not omitir_errores:
            raise UserError(self.env._(
                "El archivo tiene filas con error; corrígelo y vuelve a "
                "cargarlo, o marca «Importar sin las filas con error».\n\n%s")
                % "\n".join(self._errores_txt(con_error)))
        filas = [(n, f) for n, f in filas if n not in con_error]

        Visita = self.env["atc.visita"]
        Contratista = self.env["atc.contratista"]
        Vendedor = self.env["atc.vendedor"]
        Servicio = self.env["atc.tipo.servicio"]
        Zona = self.env["atc.zona"]
        zonas = sorted(
            ((re.sub(r"[^A-Z0-9]+", " ", _norm(z.name)).strip(), z)
             for z in Zona.search([])),
            key=lambda t: -len(t[0]))
        zonas_por_nombre = {_norm(z.name): z for _c, z in zonas}
        cache_c, cache_s, cache_f, cache_v = {}, {}, {}, {}

        abiertas = Visita.search([("tipo_orden", "=", "instalacion"),
                                  ("estado", "in", ESTADOS_ABIERTOS)])
        por_cedula = {}
        for v in abiertas:
            if v.cedula:
                por_cedula.setdefault(v.cedula, v)
        # Instalaciones ya cerradas o archivadas: la más reciente por cédula.
        # Si siguen en el reporte, las realizadas/archivadas se omiten y las
        # no realizadas se reabren, en vez de crear una orden duplicada.
        cedulas = {f["cedula"] for _n, f in filas if f["cedula"]}
        cerradas = {}
        for v in Visita.with_context(active_test=False).search(
                [("tipo_orden", "=", "instalacion"),
                 ("cedula", "in", list(cedulas - set(por_cedula))),
                 "|", ("estado", "not in", ESTADOS_ABIERTOS),
                 ("active", "=", False)],
                order="id desc"):
            cerradas.setdefault(v.cedula, v)

        actualizadas = sin_cambios = con_zona = 0
        vistas, errores, repetidas = set(), [], []
        omitidas, reabiertas, no_reabiertas, otro_contratista = [], [], [], []
        # Export completo (con ESTATUS): finalizadas ya cerradas en Odoo,
        # canceladas que no se importan y estatus que no cuadran con Odoo
        ya_realizadas = canceladas = 0
        discrepancias = []   # [(nº fila, orden, estatus en el sistema)]
        nuevas, procesadas = [], set()   # [(nº fila, vals, teléfono)]
        solo_canceladas = set()   # cédulas nuevas que solo trajeron canceladas
        for n, fila in filas:
            cedula = fila["cedula"]
            # Cédula repetida (export detallado: una fila por artículo; o el
            # mismo cliente dos veces): solo cuenta la primera fila. Se
            # comprueba antes de todo para que una repetida no cree un
            # duplicado de una orden ya tratada.
            if cedula and cedula in procesadas and not (
                    cedula in solo_canceladas and fila["estatus"] != "cancelada"):
                repetidas.append(n)
                continue
            if cedula:
                procesadas.add(cedula)
                solo_canceladas.discard(cedula)
            contratista = self._buscar_o_crear(
                Contratista, fila["contratista"], cache_c)
            servicio = self._buscar_o_crear(
                Servicio, fila["tipo_servicio"], cache_s)
            franquicia = self._buscar_franquicia(fila["franquicia"], cache_f)
            vendedor = self._buscar_o_crear(
                Vendedor, fila["vendedor"], cache_v)
            vals = {
                "tipo_servicio_id": servicio.id,
                "contratista_id": contratista.id,
                "vendedor_id": vendedor.id,
                "franquicia_id": franquicia.id,
                "telefono_movil": fila["telefono_movil"],
                "telefono_fijo": fila["telefono_fijo"],
                "direccion": fila["direccion"] or False,
            }
            existente = por_cedula.get(cedula) if cedula else None
            if existente:
                vistas.add(existente.id)
                # Solo columnas mapeadas y con valor: una celda vacía o una
                # columna que no viene no borra lo que ya tiene la orden.
                cambios = {}
                for dato, campo in ACTUALIZABLES.items():
                    nuevo = vals[campo]
                    if dato not in mapeo or not nuevo:
                        continue
                    actual = existente[campo]
                    if (actual.id if hasattr(actual, "id") else actual) == nuevo:
                        continue
                    if campo == "contratista_id" and actual:
                        # El contratista asignado en Odoo manda (ATC pudo
                        # reasignarla y el sistema no estar al día)
                        otro_contratista.append((n, existente, contratista))
                        continue
                    cambios[campo] = nuevo
                if fila["estatus"] != "abierta":
                    # No se cierra sola: el contratista o ATC deben cerrarla
                    discrepancias.append((n, existente, fila["estatus"]))
                if cambios:
                    existente.write(cambios)
                    actualizadas += 1
                else:
                    sin_cambios += 1
                continue
            cerrada = cerradas.get(cedula) if cedula else None
            if cerrada and fila["estatus"] != "abierta":
                # Cerrada en el sistema y en Odoo: no se reabre nunca
                if vendedor and not cerrada.vendedor_id:
                    cerrada.vendedor_id = vendedor
                if fila["estatus"] == "cancelada":
                    canceladas += 1
                elif cerrada.active and cerrada.estado in ("realizada", "garantia"):
                    ya_realizadas += 1
                else:
                    discrepancias.append((n, cerrada, fila["estatus"]))
                del cerradas[cedula]
                continue
            if cerrada:
                fecha_asig = self._fecha_asignacion(fila)
                cierre = cerrada.fecha_no_realizada
                reasignada = not cierre or fecha_asig > cierre.date()
                if cerrada.active and cerrada.estado == "no_realizada" \
                        and not reasignada:
                    no_reabiertas.append((n, cerrada, fecha_asig))
                elif cerrada.active and cerrada.estado == "no_realizada":
                    cambios = {campo: vals[campo]
                               for dato, campo in ACTUALIZABLES.items()
                               if dato in mapeo and vals[campo]}
                    cambios["estado"] = "asignada" if (
                        cambios.get("contratista_id") or cerrada.contratista_id
                        or cerrada.tecnico_asignado_ids) else "pendiente"
                    cambios["fecha_asignacion"] = fecha_asig
                    cerrada.write(cambios)
                    cerrada.message_post(body=self.env._(
                        "Reabierta al importar el reporte del %(r)s: el cliente "
                        "sigue pendiente de instalación y el reporte la asigna "
                        "el %(a)s, después de marcarse no realizada.",
                        r=fields.Date.to_string(self.fecha_reporte),
                        a=fields.Date.to_string(fecha_asig)))
                    reabiertas.append(cerrada)
                else:
                    omitidas.append((n, cerrada))
                # Aunque no se toque, se completa el vendedor si no lo tenía
                # (órdenes importadas antes de leer esa columna)
                if vendedor and not cerrada.vendedor_id:
                    cerrada.vendedor_id = vendedor
                del cerradas[cedula]
                continue
            if fila["estatus"] == "cancelada":
                # Nunca se instaló: no se crea; una fila posterior del mismo
                # cliente (otra orden) sí puede crearla
                canceladas += 1
                if cedula:
                    solo_canceladas.add(cedula)
                continue

            fecha_asig = self._fecha_asignacion(fila)
            zona = zonas_por_nombre.get(_norm(fila["zona"]), Zona) \
                if fila["zona"] else Zona
            if not zona:
                zona = self._zona_por_direccion(fila["direccion"], zonas)
            # Una zona de otra franquicia no se asigna (la detección por
            # dirección puede confundirse con nombres repetidos)
            if zona.franquicia_id and franquicia and zona.franquicia_id != franquicia:
                zona = Zona
            if not franquicia:
                # Sin franquicia en el archivo, la toma de la zona
                vals.pop("franquicia_id")
            con_zona += bool(zona)
            vals.update({
                "tipo_orden": "instalacion",
                "cedula": cedula or False,
                "nombre": fila["nombre"],
                "fecha": fecha_asig,
                "fecha_asignacion": fecha_asig,
                "zona_id": zona.id,
                "tecnico_id": False,
                "estado": "asignada" if contratista else "pendiente",
            })
            if fila["estatus"] == "finalizada":
                # Historial: ya instalada en el sistema. Sin serial ni
                # precinto (no vienen en el reporte); sin serial tampoco la
                # vigila la garantía automática.
                fecha_fin = fila["fecha_finalizacion"] or fecha_asig
                vals.update({
                    "estado": "realizada",
                    "fecha_instalacion": fecha_fin,
                    "descripcion": self.env._(
                        "<p>Finalizada en el sistema el %(f)s; importada como "
                        "realizada del reporte del %(r)s.</p>",
                        f=fecha_fin.strftime("%d/%m/%Y"),
                        r=self.fecha_reporte.strftime("%d/%m/%Y")),
                })
            nuevas.append((n, vals, fila["telefono_movil"]))

        creadas = self._crear_en_lote(nuevas, errores)
        n_realizadas = sum(o.estado == "realizada" for o, _t in creadas)
        for orden, telefono in creadas:
            partner = orden.cliente_id.sudo()
            if partner and not partner.phone and telefono:
                partner.phone = telefono

        fuera = len(abiertas.filtered(lambda v: v.id not in vistas))
        lineas = [
            self.env._("• Filas con datos: %s") % (len(filas) + len(con_error)),
            self.env._("• Instalaciones creadas: %(n)s (con zona: %(z)s)",
                       n=len(creadas), z=con_zona),
        ]
        if "estatus" in mapeo:
            lineas += [
                self.env._("    abiertas: %s") % (len(creadas) - n_realizadas),
                self.env._("    como realizadas (finalizadas en el sistema): %s")
                % n_realizadas,
                self.env._("• Finalizadas en el sistema que ya están realizadas "
                           "en Odoo: %s") % ya_realizadas,
                self.env._("• Canceladas en el sistema (no se importan): %s")
                % canceladas,
            ]
        lineas += [
            self.env._("• Actualizadas (contratista, vendedor, servicio, "
                       "teléfonos, dirección o franquicia): %s") % actualizadas,
            self.env._("• Sin cambios: %s") % sin_cambios,
            self.env._("• Contratistas en el archivo: %s") % len(cache_c),
            self.env._("• Franquicias en el archivo: %s") % len(cache_f),
            self.env._("• Vendedores en el archivo: %s") % len(cache_v),
            self.env._("• Instalaciones abiertas que NO vienen en este "
                       "archivo: %s") % fuera,
        ]
        if discrepancias:
            etiquetas = {"finalizada": self.env._("FINALIZADA"),
                         "cancelada": self.env._("CANCELADA")}
            estados = dict(self.env["atc.visita"]._fields["estado"]
                           ._description_selection(self.env))
            lineas.append(self.env._(
                "• Cerradas en el sistema pero no en Odoo (no se cambiaron; "
                "revisar): %s") % len(discrepancias))
            lineas += [
                "    " + self.env._(
                    "Fila %(n)s: %(nombre)s (%(cedula)s) — sistema: %(sis)s; "
                    "Odoo: %(odoo)s (%(contratista)s)", n=n, nombre=v.nombre,
                    cedula=v.cedula, sis=etiquetas[est],
                    odoo=estados[v.estado] + (
                        "" if v.active else self.env._(", archivada")),
                    contratista=v.contratista_id.name or "—")
                for n, v, est in discrepancias]
        if otro_contratista:
            lineas.append(self.env._(
                "• Abiertas con otro contratista en el reporte (se mantiene el "
                "de Odoo; revisar): %s") % len(otro_contratista))
            lineas += [
                "    " + self.env._(
                    "Fila %(n)s: %(nombre)s (%(cedula)s) — Odoo: %(odoo)s; "
                    "reporte: %(rep)s", n=n, nombre=v.nombre, cedula=v.cedula,
                    odoo=v.contratista_id.name, rep=c.name)
                for n, v, c in otro_contratista]
        if no_reabiertas:
            lineas.append(self.env._(
                "• No realizadas que siguen en el reporte, NO reabiertas (el "
                "reporte no las asigna después del cierre; el sistema puede no "
                "estar al día): %s") % len(no_reabiertas))
            lineas += [
                "    " + self.env._(
                    "Fila %(n)s: %(nombre)s (%(cedula)s) — no realizada el "
                    "%(cierre)s por %(contratista)s: %(motivos)s; asignada en "
                    "el reporte el %(asig)s", n=n, nombre=v.nombre,
                    cedula=v.cedula,
                    cierre=fields.Date.to_string(v.fecha_no_realizada.date()),
                    contratista=v.contratista_id.name or "—",
                    motivos=", ".join(v.motivo_no_realizada_ids.mapped("name")) or "—",
                    asig=fields.Date.to_string(fecha))
                for n, v, fecha in no_reabiertas]
        if reabiertas:
            lineas.append(self.env._(
                "• No realizadas reasignadas en el reporte después del cierre, "
                "reabiertas: %s") % len(reabiertas))
            lineas += ["    %s (%s)" % (v.nombre, v.cedula) for v in reabiertas]
        if omitidas:
            lineas.append(self.env._(
                "• Ya realizadas o archivadas en Odoo, omitidas (revisar si "
                "hay que reinstalar): %s") % len(omitidas))
            lineas += [
                "    " + self.env._(
                    "Fila %(n)s: %(nombre)s (%(cedula)s) — %(estado)s",
                    n=n, nombre=v.nombre, cedula=v.cedula,
                    estado=dict(v._fields["estado"]._description_selection(
                        self.env))[v.estado]
                    + ("" if v.active else self.env._(", archivada")))
                for n, v in omitidas]
        if repetidas:
            lineas.append(self.env._(
                "• Cédula repetida en el archivo (se usó la primera), "
                "filas: %s") % ", ".join(map(str, repetidas)))
        if con_error:
            lineas.append(self.env._(
                "• Filas con error en el archivo (NO importadas, corregir y "
                "volver a cargar): %s") % len(con_error))
            lineas += ["    %s" % e for e in self._errores_txt(con_error)]
        if errores:
            lineas.append(self.env._("• Filas con error (no importadas): %s")
                          % len(errores))
            lineas += ["    %s" % e for e in errores]
        if avisos:
            lineas.append(self.env._("• Avisos:"))
            lineas += ["    %s" % a for a in avisos]
        return "\n".join(lineas)

    def _crear_en_lote(self, nuevas, errores):
        """Crea todas de una vez; si alguna falla, las crea una por una para
        saltar solo las filas con error."""
        # atc_cierre_automatico: las finalizadas en el sistema entran como
        # realizadas sin serial, técnico ni precinto
        Visita = self.env["atc.visita"].with_context(atc_cierre_automatico=True)
        if not nuevas:
            return []
        try:
            with self.env.cr.savepoint():
                ordenes = Visita.create([vals for _n, vals, _t in nuevas])
            return list(zip(ordenes, [t for _n, _v, t in nuevas]))
        except Exception:
            _logger.info("Import instalaciones: el lote falló, se crea fila a fila")
        creadas = []
        for n, vals, telefono in nuevas:
            try:
                with self.env.cr.savepoint():
                    creadas.append((Visita.create(vals), telefono))
            except Exception as exc:
                msg = exc.args[0] if exc.args else str(exc)
                errores.append(self.env._("Fila %(n)s: %(msg)s", n=n, msg=msg))
        return creadas

    def _recordar_mapeo(self):
        recordado = self._mapeo_recordado()
        for linea in self.linea_ids:
            clave = _norm(linea.columna)
            if clave:
                recordado[clave] = linea.campo or ""
        self.env["ir.config_parameter"].sudo().set_param(
            PARAM_MAPEO, json.dumps(recordado, ensure_ascii=False))

    def _reabrir(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
            "name": self.env._("Importar instalaciones"),
        }

    # ------------------------------------------------------------------
    # Botones
    # ------------------------------------------------------------------
    def action_cargar(self):
        self._cargar()
        return self._reabrir()

    def action_releer(self):
        self.ensure_one()
        if self.fila_encabezado < 1:
            raise UserError(self.env._("La fila de encabezados empieza en 1."))
        self._cargar(hoja=self.hoja, indice=self.fila_encabezado - 1)
        return self._reabrir()

    def action_volver(self):
        self.linea_ids.unlink()
        self.write({"state": "input", "resultado": False, "datos": False})
        return self._reabrir()

    def action_probar(self):
        """Simula la importación y deshace todo; muestra el resumen."""
        self.ensure_one()

        class _Simulacion(Exception):
            pass

        try:
            with self.env.cr.savepoint():
                # En la prueba se muestran las filas con error en el resumen
                # en vez de cortar.
                resumen = self._importar(omitir_errores=True)
                raise _Simulacion()
        except _Simulacion:
            pass
        self.env.invalidate_all()
        self.resultado = self.env._(
            "PRUEBA (no se guardó nada):\n\n%s") % resumen
        return self._reabrir()

    def action_importar(self):
        self.ensure_one()
        resumen = self._importar(omitir_errores=self.omitir_errores)
        self._recordar_mapeo()
        _logger.info("Import instalaciones: %s", resumen.replace("\n", " | "))
        self.write({"state": "done",
                    "resultado": self.env._("Importación completada.\n\n")
                    + resumen})
        return self._reabrir()


class AtcInstalacionImportColumna(models.TransientModel):
    _name = "atc.instalacion.import.columna"
    _description = "Columna del archivo a importar"
    _order = "indice"

    wizard_id = fields.Many2one(
        "atc.instalacion.import", required=True, ondelete="cascade")
    indice = fields.Integer(readonly=True)
    columna = fields.Char(string="Columna del archivo", readonly=True)
    ejemplo = fields.Char(string="Valores de ejemplo", readonly=True)
    campo = fields.Selection(
        CAMPOS, string="Importar como",
        help="Vacío = no importar esta columna.")
